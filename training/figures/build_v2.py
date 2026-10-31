"""Produce SmartAgro_Manuscript_v2.docx from the current manuscript.

Applies, in order: the Part A author fix, the Part B/C content items that are
still outstanding, then the Part C formatting requirements.
"""

import copy
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image

SRC = sys.argv[1]
DST = sys.argv[2]
FIGDIR = sys.argv[3]

doc = Document(SRC)
body = doc.element.body
log = []


def note(msg):
    log.append(msg)
    print(msg)


# ---------------------------------------------------------------- helpers
def blocks():
    """Body-level paragraphs and tables, in document order."""
    out = []
    for ch in body.iterchildren():
        if ch.tag == qn("w:p"):
            out.append(Paragraph(ch, doc))
        elif ch.tag == qn("w:tbl"):
            out.append(Table(ch, doc))
    return out


def paras():
    return [b for b in blocks() if isinstance(b, Paragraph)]


def find_para(prefix, start=0):
    ps = paras()
    for i, p in enumerate(ps):
        if i >= start and p.text.strip().startswith(prefix):
            return p
    return None


def set_text(p, text):
    """Replace a paragraph's text, keeping the formatting of its first run."""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)


BODY_TEMPLATE = None


def body_para(text):
    """A new justified body paragraph matching the manuscript's body style."""
    global BODY_TEMPLATE
    if BODY_TEMPLATE is None:
        for p in paras():
            if (p.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY and p.runs
                    and len(p.text) > 200 and not p.runs[0].bold):
                BODY_TEMPLATE = p
                break
    new = copy.deepcopy(BODY_TEMPLATE._p)
    np_ = Paragraph(new, doc)
    set_text(np_, text)
    for r in np_.runs:
        r.bold = None
        r.italic = None
    return np_


def insert_after(el, para):
    el.addnext(para._p)
    return para._p


# ============================================================ PART A
# A2: Yash Chaubey becomes author 2, in the byline and in CRediT.
BYLINE_OLD = ("Utkarsh Persai, Venkata Hrishi Samay Jetty, Devansh Rathore, "
              "Yash Chaubey, Syamasudha Veeragandham*")
BYLINE_NEW = ("Utkarsh Persai, Yash Chaubey, Venkata Hrishi Samay Jetty, "
              "Devansh Rathore, Syamasudha Veeragandham*")
p = find_para("Utkarsh Persai,")
assert p is not None and p.text.strip() == BYLINE_OLD, repr(p.text if p else None)
set_text(p, BYLINE_NEW)
note("A2 byline reordered")

CREDIT_NEW = (
    "Utkarsh Persai: Conceptualization, Methodology, Software, Formal analysis, "
    "Investigation, Data curation, Visualization, Writing \u2013 original draft. "
    "Yash Chaubey: Data curation, Validation, Visualization, Writing \u2013 review "
    "and editing. "
    "Venkata Hrishi Samay Jetty: Software, Validation, Investigation, Writing \u2013 "
    "review and editing. "
    "Devansh Rathore: Hardware, Validation, Investigation, Resources. "
    "Syamasudha Veeragandham: Supervision, Project administration, "
    "Conceptualization, Writing \u2013 review and editing.")
p = find_para("Utkarsh Persai: Conceptualization")
assert p is not None
set_text(p, CREDIT_NEW)
note("A2 CRediT reordered to match")


# A2b: superscript affiliation markers on the byline. The paper carries one
# affiliation, so every author takes marker 1; the corresponding author also
# takes the asterisk.
AUTHORS = [
    ("Utkarsh Persai", "1"),
    ("Yash Chaubey", "1"),
    ("Venkata Hrishi Samay Jetty", "1"),
    ("Devansh Rathore", "1"),
    ("Syamasudha Veeragandham", "1,*"),
]


def rebuild_byline(p, authors):
    """Rewrite the byline as name / superscript-marker pairs."""
    template = p.runs[0]._element
    base = template.find(qn("w:rPr"))
    for r in list(p.runs):
        r._element.getparent().remove(r._element)

    def add(text, sup=False):
        run = p.add_run(text)
        if base is not None:
            run._element.insert(0, copy.deepcopy(base))
        run.bold = None
        run.italic = None
        if sup:
            run.font.superscript = True
        return run

    for i, (name, mark) in enumerate(authors):
        if i:
            add(", ")
        add(name)
        add(mark, sup=True)


p = find_para("Utkarsh Persai, Yash Chaubey")
assert p is not None
rebuild_byline(p, AUTHORS)
note("A2b byline numbered: " +
     ", ".join("%s(%s)" % (n.split()[-1], m) for n, m in AUTHORS))

# matching marker in front of the affiliation line
p = find_para("School of Computer Science and Engineering")
assert p is not None
first = p.runs[0]
marker = p.insert_paragraph_before("")._p  # placeholder, removed below
marker.getparent().remove(marker)
new_run = p.add_run("1")
new_run.font.superscript = True
rpr = first._element.find(qn("w:rPr"))
if rpr is not None:
    new_run._element.insert(0, copy.deepcopy(rpr))
new_run.font.superscript = True
p._p.insert(list(p._p).index(first._element), new_run._element)
note("A2b affiliation marked with superscript 1")


# ============================================================ PART C5
# The one real footnote carries the fusion disclosure, which must not be lost,
# so its text is promoted into the body and the reference removed.
FOOT_TEXT = ("In the deployed code, three normalised sensor values are appended "
             "to the image features at inference time. The saved classifier was "
             "trained on image features only, so the inference code trims the "
             "vector back to its training length and the sensor values are "
             "dropped. Every prediction from the deployed system is therefore "
             "image-only, which is also what is evaluated in this paper.")
p = find_para("A separate point concerns the multimodal fusion stage")
assert p is not None
for r in list(p.runs):
    if r._element.findall(qn("w:footnoteReference")):
        r._element.getparent().remove(r._element)
insert_after(p._p, body_para(FOOT_TEXT))
note("C5 footnote removed; fusion disclosure promoted to body text")

# strip the footnote itself so the file carries none
fn_part = None
for rel in doc.part.rels.values():
    if rel.reltype.endswith("/footnotes"):
        fn_part = rel.target_part
if fn_part is not None:
    from lxml import etree
    froot = etree.fromstring(fn_part.blob)
    dropped = 0
    for fn in list(froot):
        if fn.get(qn("w:type")) is None:
            froot.remove(fn)
            dropped += 1
    fn_part._blob = etree.tostring(froot, xml_declaration=True,
                                   encoding="UTF-8", standalone=True)
    note("C5 footnotes.xml reduced to separators only (%d dropped)" % dropped)


# ============================================================ PART C6
# A "where:" list after every equation. Values are taken from
# hardware/smart_agri/smart_agri.ino, training/source_disjoint_eval.py,
# training/train_soil_model.py and results/probe_check_simulation.json.
WHERE = {
    "M = clamp(": [
        "M: soil moisture reported by the firmware, % (0\u2013100);",
        "r: raw ADC reading on GPIO 34, dimensionless (0\u20134095, 12-bit);",
        "A: reading with the probe in dry air, 2590 (firmware MOISTURE_AIR);",
        "W: reading with the probe in water, 1500 (firmware MOISTURE_WATER);",
        "clamp(x, 0, 100): limits x to 0\u2013100, implemented as Arduino "
        "constrain(); the mapping itself uses integer arithmetic.",
    ],
    "u_t = 1 if": [
        "u\u209c: pump state at sample t (0 = off, 1 = on);",
        "M\u209c: moisture at sample t, % (0\u2013100);",
        "T_D: dry threshold, 30% (firmware MOISTURE_DRY);",
        "T_W: wet threshold, 60% (firmware MOISTURE_WET);",
        "t: sample index, one per second (delay(1000) in the loop).",
    ],
    "x = [ \u03c6(I)": [
        "I: input RGB image, 224 \u00d7 224 \u00d7 3, pixel values 0\u2013255;",
        "\u03c6(I): frozen backbone embedding after global average pooling; "
        "1,280 values for EfficientNetV2-B0 and MobileNetV2, 2,048 for ResNet50;",
        "c(I): colour descriptor, 102 values \u2014 a 32-bin histogram for each "
        "RGB channel, each histogram normalised to sum to 1 (each bin 0\u20131), "
        "plus the mean and standard deviation of each channel divided by 255 "
        "(0\u20131); used for soil type only;",
        "x: concatenated feature vector, 1,382 values for the two "
        "1,280-dimensional backbones and 2,150 for ResNet50;",
        "[\u00b7 , \u00b7]: concatenation.",
    ],
    "f_kl(x) = sign": [
        "f_kl(x): decision for the pair of classes k and l; scikit-learn SVC "
        "votes over all K(K \u2212 1)/2 pairs, with K = 5 for soil type and "
        "K = 4 for soil condition;",
        "x\u1d62: support vectors drawn from the training folds;",
        "y\u1d62: pairwise labels, \u00b11;",
        "\u03b1\u1d62: learned dual coefficients, |\u03b1\u1d62| \u2264 C with "
        "C = 10 (fixed a priori);",
        "\u03b3: RBF kernel width, gamma='scale' = 1/(n_features \u00b7 Var(X)) "
        "computed on each standardised training fold (> 0);",
        "b_kl: bias for that pair;",
        "\u2016\u00b7\u2016: Euclidean norm.",
    ],
    "\u0394 = Acc_aug": [
        "\u0394: protocol gap, percentage points (\u2212100 to +100; positive "
        "means the augmentation-level protocol scores higher);",
        "Acc_aug: accuracy under stratified 5-fold cross-validation over the "
        "augmented samples, % (0\u2013100);",
        "Acc_disjoint: accuracy under leave-one-photograph-out "
        "cross-validation, % (0\u2013100).",
    ],
    "CI_W = [": [
        "CI_W: Wilson score interval for a proportion (0\u20131);",
        "p\u0302: observed accuracy as a proportion (0\u20131);",
        "n: number of samples (4,194 for soil type, 294 for soil condition);",
        "z: standard normal quantile, 1.96 for a 95% interval.",
    ],
    "p = #{": [
        "p: two-sided exact p-value (0\u20131); the smallest value attainable "
        "with n photographs is 2^(1\u2212n);",
        "n: number of source photographs (16 for soil type, 14 for soil "
        "condition), so 65,536 and 16,384 sign assignments are enumerated;",
        "\u03c3: a vector of signs in {\u22121, +1}\u207f;",
        "d_s: difference in accuracy on photograph s between the two "
        "protocols (\u22121 to +1);",
        "mean(\u00b7): arithmetic mean over the n photographs.",
    ],
    "g = max_t": [
        "M\u2080: mean moisture over the 5 s before the pulse, % (0\u2013100);",
        "M(t): moisture after the pulse, sampled at 10 Hz for 90 s, % "
        "(0\u2013100);",
        "V: delivered volume, mL (> 0); 50 mL in the simulation;",
        "g: steepest rise, percentage points per second (> 0); flagged above "
        "g_max = 8.0;",
        "\u03c0: peak rise per unit volume, percentage points per mL (> 0); "
        "flagged below \u03c0_min = 0.239;",
        "\u03c4: onset delay before the moisture begins to rise, s (\u2265 0); "
        "flagged below \u03c4_min = 3.6.",
    ],
}

for key, items in WHERE.items():
    ps = paras()
    p = None
    idx = None
    for i, q in enumerate(ps):
        if q.text.strip().startswith(key):
            p, idx = q, i
            break
    if p is None:
        note("  !! equation not found: %r" % key)
        continue
    anchor = p._p
    # Eqs (2) and (6) sit inside a sentence that finishes in the paragraph
    # below; dropping a "where:" list between the two halves would break the
    # sentence, so in that case the list goes after the completing clause.
    nxt = ps[idx + 1] if idx + 1 < len(ps) else None
    if nxt is not None:
        t = nxt.text.strip()
        if t and t[0].islower():
            anchor = nxt._p
            note("  (where: placed after the clause completing %s)" % key.strip()[:14])
    head = body_para("where:")
    insert_after(anchor, head)
    anchor = head._p
    for it in items:
        ip = body_para("    " + it)
        insert_after(anchor, ip)
        anchor = ip._p
note("C6 'where:' lists added after %d equations" % len(WHERE))
